from openpyxl import load_workbook

from job_agent.export import workbook_path, write_workbook
from job_agent.report import Report


def job(job_id: str, bucket: str, score: float | None) -> dict:
    return {"Job ID": job_id, "_bucket": bucket, "Score": score, "Title": f"Job {job_id}",
            "Company": "Acme", "New": "", "Apply URL": f"https://x/jobs/{job_id}",
            "Work mode": "remote", "Realistic (₹ LPA)": "", "Listed (₹ LPA)": "",
            "Reject reason": "sales" if bucket == "rejected" else ""}  # fmt: skip


REPORT = Report(
    jobs=[job("1", "remote_foreign", 90), job("2", "india_onsite", 50), job("3", "rejected", None)],
    companies=[{"Company": "Acme", "Website": "https://acme.dev"}],
)


def test_tabs_and_routing(tmp_path):
    path = write_workbook(REPORT, workbook_path(tmp_path, "2026-10-03"), threshold=75)
    wb = load_workbook(path)
    assert path.name == "jobs_2026-10-03.xlsx"
    ids = lambda sheet: [r[-1] for r in wb[sheet].iter_rows(min_row=3, values_only=True)]  # noqa: E731
    assert wb["🔥 Top Picks"].cell(3, 5).value == "Job 1"
    assert ids("🗑️ Rejected") == ["3"]
    assert wb["🏢 India · Onsite"].cell(3, 5).value == "Job 2"


def test_dashboard_top10_only_scored_and_pending_count(tmp_path):
    report = Report(jobs=[job("1", "needs_review", None), job("2", "remote_india", 70)], pending=5)
    wb = load_workbook(write_workbook(report, tmp_path / "x.xlsx", threshold=75))
    cells = [c for row in wb["📊 Dashboard"].iter_rows(values_only=True) for c in row if c]
    assert "⏳ Not processed yet" in cells and 5 in cells
    assert any("Job 2" in str(c) for c in cells) and not any("Job 1" in str(c) for c in cells)
