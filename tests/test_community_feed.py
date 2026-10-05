from openpyxl import load_workbook

import job_agent.community.feed as feed
from job_agent.community.digest import digest_text, write_workbook
from job_agent.community.events import Event
from job_agent.community.models import Posting
from job_agent.community.store import CommunityStore
from job_agent.settings import Settings


def p(source, company, title, location, kind_pay=None, **extra):
    return Posting(source, company, title[:3] + company, title, f"https://x/{company}",
                   location=location, extra={"pay": kind_pay, **extra})  # fmt: skip


RAW = [
    (p("lever", "CRED", "SDE I", "Bengaluru"), "unicorn"),
    (p("lever", "CRED", "Senior Software Engineer", "Bengaluru"), "unicorn"),
    (p("greenhouse", "Stripe", "Software Engineer", "Singapore"), "unicorn"),
    (p("greenhouse", "Stripe", "Software Engineering Intern", "Bengaluru"), "unicorn"),
    (p("unstop", "Acme Edu", "Python Internship", "Remote, India", internship=True), "other"),
    (p("unstop", "DataCo", "Data Science Internship", "Remote, India", "₹15,000–20,000",
       internship=True), "other"),
    (p("adzuna", "Zomato", "Account Manager", "Gurgaon"), "other"),
]  # fmt: skip
EVENTS = [Event("devfolio", "1", "Hackathon", "HackX", "IIT H", "In person", "Hyderabad",
                "2026-10-20", "2026-10-21", "2026-10-15", "", "https://hackx.devfolio.co"),
          Event("gdg", "2", "Meetup / workshop", "Gemini Study Jam", "GDG on Campus X - Pune, India",
                "Online", "on Campus X University - Pune, India", "2026-10-09", None, None, "",
                "https://gdg.community.dev/e/1/")]  # fmt: skip


def test_refresh_curates_and_counts(tmp_path, monkeypatch):
    monkeypatch.setattr(feed, "gather", lambda settings: (RAW, {}))
    monkeypatch.setattr(feed, "fetch_events", lambda client: (EVENTS, {}))
    with CommunityStore(tmp_path / "c.db") as store:
        stats = feed.refresh(Settings(), store)
        assert stats.kept == {"job": 1, "internship": 2}
        assert stats.dropped["Senior role"] == 1 and stats.dropped["Not in India"] == 1
        assert stats.dropped["Internship without stipend"] == 1
        assert stats.dropped["Not a tech role"] == 1 and stats.new["event"] == 2
        rows = store.live_postings(stats.run_at)
        stripe = next(r for r in rows if r["company"] == "Stripe")
        assert stripe["kind"] == "internship" and stripe["pay"] == "Paid (amount not listed)"
        gdg = next(e for e in store.live_events(stats.run_at, "2026-10-05") if e["source"] == "gdg")
        assert gdg["city"] == "Pune (X University)"

        jobs = [r for r in rows if r["kind"] == "job"]
        interns = [r for r in rows if r["kind"] == "internship"]
        events = store.live_events(stats.run_at, "2026-10-05")
        today = stats.run_at[:10]
        text = digest_text(jobs, interns, events, today)
        assert "SDE I" in text and "₹15,000–20,000" in text and "HackX" in text
        path = write_workbook(tmp_path / "feed.xlsx", jobs, interns, events, today)
        wb = load_workbook(path)
        assert wb.sheetnames == ["📊 Today", "💼 Jobs", "🎓 Internships", "🏆 Hackathons",
                                 "🎤 Tech events", "🔍 Check experience"]  # fmt: skip
        assert wb["🎓 Internships"].cell(2, 7).value in (
            "Paid (amount not listed)",
            "₹15,000–20,000",
        )


def test_listing_only_postings_get_their_description(tmp_path, monkeypatch):
    texts = {"wd1": "We need 5+ years of experience in C++.", "wd2": "Open to fresher graduates."}
    calls = []

    def describe(posting, client):
        calls.append(posting.external_id)
        return texts[posting.external_id]

    raw = [(Posting("workday", "NVIDIA", "wd1", "Software Engineer", "https://x/1",
                    location="Bengaluru, India"), "big_tech"),
           (Posting("workday", "NVIDIA", "wd2", "Software Engineer", "https://x/2",
                    location="Pune, India"), "big_tech")]  # fmt: skip
    monkeypatch.setattr(feed, "gather", lambda settings: (raw, {}))
    monkeypatch.setattr(feed, "fetch_events", lambda client: ([], {}))
    monkeypatch.setitem(feed.DESCRIBERS, "workday", describe)
    with CommunityStore(tmp_path / "c.db") as store:
        stats = feed.refresh(Settings(), store)
        assert stats.described == 2 and stats.dropped["Needs 5+ years"] == 1
        [row] = store.live_postings(stats.run_at)
        assert row["level"] == "Entry level"
        for p_, _ in raw:
            p_.description = ""
        feed.refresh(Settings(), store)  # second run reads the cache
        assert calls == ["wd1", "wd2"]


def test_hack2skill_keeps_open_events_only():
    import httpx

    from job_agent.community.events import hack2skill

    data = {"data": {"flagshipEvents": [
        {"_id": "1", "title": "AI Builder Cup 2026", "status": "APPROVED",
         "tags": {"mode": {"value": "HYBRID"}}, "registrationEnd": "2026-10-11T05:00:00Z",
         "customEventUrl": "https://aibuildercup.com"},
        {"_id": "2", "title": "Old Hackathon", "status": "APPROVED",
         "tags": {"mode": {"value": "VIRTUAL"}}, "registrationEnd": "2026-01-01T00:00:00Z"},
    ], "communityEvents": [
        {"_id": "3", "title": "Build with AI: Bootcamps", "eventUrl": "bootcamp",
         "tags": {"mode": {"value": "VIRTUAL"}}, "registrationEnd": "2026-11-30T00:00:00Z"},
    ]}}  # fmt: skip
    c = httpx.Client(transport=httpx.MockTransport(lambda r: httpx.Response(200, json=data)))
    cup, camp = hack2skill(c, today="2026-10-05")
    assert (cup.kind, cup.mode, cup.deadline, cup.url) == (
        "Hackathon",
        "Hybrid",
        "2026-10-11",
        "https://aibuildercup.com",
    )
    assert camp.kind == "Meetup / workshop" and camp.mode == "Online"
    assert camp.url == "https://hack2skill.com/event/bootcamp/"


def big(key, company, tier="big_tech", first_seen="2026-10-05T08:00", level="Entry level",
        kind="job", posted="2026-10-01"):  # fmt: skip
    return {"key": key, "company": company, "title": f"Engineer {key}", "tier": tier,
            "first_seen": first_seen, "level": level, "kind": kind, "posted_at": posted,
            "url": "https://x", "location": "Pune", "pay": None, "source": "x"}  # fmt: skip


def test_big_picks_prefer_new_then_unfeatured_one_per_company():
    from job_agent.community.digest import big_picks

    since = "2026-10-05T08:00"
    rows = [
        big("a", "Startup", tier="startup"),
        big("b", "Cisco", first_seen="2026-09-01T08:00"),
        big("c", "Cisco"), big("d", "Cisco"),
        big("e", "Visa", tier="mnc", level="Not specified"),
        big("f", "Intel", first_seen="2026-09-01T08:00"),
    ]  # fmt: skip
    assert [r["key"] for r in big_picks(rows, [], since)] == ["c", "e", "f"]  # one per company
    # Nothing new: the ones featured lately go last.
    old = [{**r, "first_seen": "2026-09-01T08:00"} for r in rows]
    assert [r["key"] for r in big_picks(old, [], since, recent={"b", "c"})][:2] == ["d", "f"]


def test_todays_picks_are_stable_and_rotate(tmp_path):
    from job_agent.community.service import todays_picks

    rows = [big(k, f"Co {k}", first_seen="2026-09-01T08:00") for k in "abcdef"]
    with CommunityStore(tmp_path / "c.db") as store:
        store.set_meta("last_refresh", "2026-10-05T08:00")
        first = [r["key"] for r in todays_picks(store, rows, [], "2026-10-05")]
        assert first == [r["key"] for r in todays_picks(store, rows, [], "2026-10-05")]
        nxt = [r["key"] for r in todays_picks(store, rows, [], "2026-10-06")]
        assert len(first) == 3 and not set(first) & set(nxt)
