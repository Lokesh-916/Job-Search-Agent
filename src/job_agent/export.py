"""Daily workbook writer (see docs/WORKBOOK.md) and the status read-back from the last one."""

from __future__ import annotations

from collections.abc import Sequence
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import xlsxwriter
from openpyxl import load_workbook

from job_agent.report import Report

STATUSES = ["To Apply", "Applied", "Interviewing", "Offer", "Rejected", "Skip"]
USER_COLUMNS = ("Status", "My Notes", "Applied On")
URL_COLUMNS = {"Apply URL", "Job URL", "Website"}


@dataclass(frozen=True)
class Band:
    name: str
    color: str
    columns: tuple[str, ...]
    collapsed: bool = False


JOB_BANDS = (
    Band("Tracking", "#1F4E79", ("Score", "Tier", "New", "Status", "My Notes", "Applied On")),
    Band("Role", "#2E75B6", ("Title", "Company", "Category", "Builds AI?", "Apply URL",
                             "Job URL", "Posted", "Age (days)", "Openings at company")),
    Band("Location", "#548235", ("Work mode", "Locations", "India eligible",
                                 "Eligibility evidence", "Timezone / hours", "Visa")),
    Band("Pay", "#BF8F00", ("Listed salary", "Listed (₹ LPA)", "Equity", "Realistic (₹ LPA)",
                            "Pay confidence", "Pay basis")),
    Band("Fit", "#7030A0", ("Verdict", "Fit score", "Why I fit", "Pitch", "Learning upside",
                            "Joining fit", "Red flags")),
    Band("Interview", "#C55A11", ("DSA risk", "DSA evidence", "Take-home / practical?",
                                  "Interview process"), collapsed=True),
    Band("Requirements", "#808080", ("Experience", "Fresher OK?", "Must-have skills",
                                     "Nice-to-have", "Tech stack", "Gaps"), collapsed=True),
    Band("Company", "#305496", ("What they do", "YC batch", "Stage", "Team size", "HQ",
                                "Employee sentiment", "Founders"), collapsed=True),
    Band("Meta", "#595959", ("First seen", "Job ID"), collapsed=True),
)  # fmt: skip
REJECTED_COLUMNS = ("Reject reason", "Title", "Company", "Category", "Work mode", "Visa",
                    "Listed salary", "Experience", "Apply URL", "Job ID")  # fmt: skip

WIDTHS = {
    "Score": 7, "Tier": 5, "New": 5, "Status": 13, "My Notes": 24, "Title": 34, "Company": 18,
    "Builds AI?": 30, "Apply URL": 9, "Job URL": 9, "Locations": 22, "Eligibility evidence": 30,
    "Pay basis": 30, "Why I fit": 45, "Pitch": 40, "Learning upside": 30, "DSA evidence": 30,
    "Interview process": 40, "Must-have skills": 30, "Tech stack": 26, "What they do": 45,
    "Founders": 35, "Reject reason": 45, "Employee sentiment": 35, "Red flags": 25, "Gaps": 28,
    "One-liner": 40, "Pros": 35, "Cons": 35, "Interview experiences": 40, "Outreach draft": 50,
    "Salary data points": 30, "Recent news": 40, "Sources": 30,
}  # fmt: skip
WRAP = {"Why I fit", "Pitch", "Pay basis", "Eligibility evidence", "DSA evidence",
        "Interview process", "What they do", "Reject reason", "Employee sentiment",
        "Pros", "Cons", "Interview experiences", "Outreach draft", "One-liner"}  # fmt: skip

TABS = (  # (sheet name, filter)
    ("🔥 Top Picks", "top"),
    ("🌍 Remote · Foreign", "remote_foreign"),
    ("🏠 Remote · India", "remote_india"),
    ("🏢 India · Onsite", "india_onsite"),
    ("✈️ Abroad · Sponsored", "abroad_sponsored"),
    ("🔎 Needs Review", "needs_review"),
)


def workbook_path(output_dir: Path, run_date: str) -> Path:
    return output_dir / f"jobs_{run_date}.xlsx"


def read_user_status(output_dir: Path) -> dict[str, dict[str, str | None]]:
    """Status / notes the user typed into the most recent workbook, keyed by Job ID."""
    books = sorted(output_dir.glob("jobs_*.xlsx"))
    if not books:
        return {}
    wb = load_workbook(books[-1], read_only=True, data_only=True)
    found: dict[str, dict[str, str | None]] = {}
    for ws in wb.worksheets:
        rows = ws.iter_rows(min_row=2, values_only=True)  # row 1 = band titles
        header = next(rows, None)
        if not header or "Job ID" not in header or "Status" not in header:
            continue
        idx = {name: header.index(name) for name in ("Job ID", *USER_COLUMNS) if name in header}
        for values in rows:
            job_id = values[idx["Job ID"]]
            if job_id is None:
                continue
            entry = {
                col.lower().replace(" ", "_").replace("my_", ""): _cell(values[i])
                for col, i in idx.items()
                if col != "Job ID"
            }
            if any(entry.values()):
                found[str(job_id)] = entry
    wb.close()
    return found


def _cell(v: Any) -> str | None:
    if v is None or v == "":
        return None
    return v.isoformat()[:10] if hasattr(v, "isoformat") else str(v)


class _Writer:
    def __init__(self, path: Path):
        self.wb = xlsxwriter.Workbook(str(path), {"strings_to_urls": False})
        self.header = self.wb.add_format(
            {"bold": True, "font_color": "white", "bg_color": "#1F4E79", "text_wrap": True}
        )
        self.wrap = self.wb.add_format({"text_wrap": True, "valign": "top"})
        self.top = self.wb.add_format({"valign": "top"})
        self.link = self.wb.add_format({"font_color": "#0563C1", "underline": 1, "valign": "top"})
        self.title = self.wb.add_format({"bold": True, "font_size": 16})
        self.bold = self.wb.add_format({"bold": True})

    def band_format(self, color: str):
        return self.wb.add_format({"bold": True, "font_color": "white", "bg_color": color,
                                   "align": "center"})  # fmt: skip

    def table(
        self,
        name: str,
        rows: Sequence[dict[str, Any]],
        bands: Sequence[Band],
        freeze_cols: int = 0,
    ) -> None:
        ws = self.wb.add_worksheet(name)
        columns = [c for b in bands for c in b.columns]
        col = 0
        for band in bands:
            last = col + len(band.columns) - 1
            fmt = self.band_format(band.color)
            if last > col:
                ws.merge_range(0, col, 0, last, band.name, fmt)
            else:
                ws.write(0, col, band.name, fmt)
            for c in range(col, last + 1):
                header = columns[c]
                opts = {"level": 1, "hidden": True} if band.collapsed else {}
                fmt = self.wrap if header in WRAP else self.top
                ws.set_column(c, c, WIDTHS.get(header, 14), fmt, opts)
            col = last + 1

        data = [[r.get(c) for c in columns] for r in rows] or [[None] * len(columns)]
        ws.add_table(1, 0, len(data) + 1, len(columns) - 1, {
            "data": data, "style": "Table Style Light 9", "name": _table_name(name),
            "columns": [{"header": c, "header_format": self.header} for c in columns],
        })  # fmt: skip
        for r, row in enumerate(rows, start=2):  # rewrite URL cells as real hyperlinks
            for c, header in enumerate(columns):
                if header in URL_COLUMNS and row.get(header):
                    ws.write_url(r, c, row[header], self.link, string="open")
        last_row = len(data) + 1
        if "Status" in columns:
            s = columns.index("Status")
            ws.data_validation(
                2, s, max(last_row, 500), s, {"validate": "list", "source": STATUSES}
            )
        for score_col in ("Score", "Fit score"):
            if score_col in columns:
                c = columns.index(score_col)
                ws.conditional_format(2, c, last_row, c, {
                    "type": "3_color_scale", "min_color": "#F8696B", "mid_color": "#FFEB84",
                    "max_color": "#63BE7B",
                })  # fmt: skip
        ws.freeze_panes(2, freeze_cols)
        ws.set_row(1, 30)

    def key_values(self, name: str, title: str, items: Sequence[tuple[str, Any]]) -> Any:
        ws = self.wb.add_worksheet(name)
        ws.write(0, 0, title, self.title)
        ws.set_column(0, 0, 28)
        ws.set_column(1, 1, 60)
        for i, (k, v) in enumerate(items, start=2):
            ws.write(i, 0, k, self.bold)
            ws.write(i, 1, v)
        return ws

    def close(self) -> None:
        self.wb.close()


def _table_name(sheet: str) -> str:
    return "T_" + "".join(ch for ch in sheet if ch.isascii() and ch.isalnum())


def write_workbook(
    report: Report,
    path: Path,
    threshold: float,
    run_info: Sequence[tuple[str, Any]] = (),
) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    live = [j for j in report.jobs if j["_bucket"] != "rejected"]
    tabs = {
        name: [j for j in live if (j["Score"] or 0) >= threshold]
        if key == "top"
        else [j for j in live if j["_bucket"] == key]
        for name, key in TABS
    }
    rejected = [j for j in report.jobs if j["_bucket"] == "rejected"]

    w = _Writer(path)
    dash = w.key_values(
        "📊 Dashboard",
        "Job-Search-Agent · daily picks",
        [
            *run_info,
            ("", ""),
            *((name, len(rows)) for name, rows in tabs.items()),
            ("🏭 Companies", len(report.companies)),
            ("🗑️ Rejected", len(rejected)),
            ("🆕 New today", sum(1 for j in live if j["New"])),
        ],
    )
    top10 = sorted(live, key=lambda j: -(j["Score"] or 0))[:10]
    start = len(run_info) + len(tabs) + 7
    dash.write(start, 0, "Top 10", w.title)
    for i, j in enumerate(top10, start=start + 1):
        dash.write(i, 0, j["Score"])
        dash.write_url(i, 1, j["Apply URL"], w.link, string=f"{j['Title']} · {j['Company']}")
        dash.write(i, 2, j["Realistic (₹ LPA)"] or j["Listed (₹ LPA)"])
        dash.write(i, 3, j["Work mode"])
    dash.set_column(2, 3, 16)

    for name, rows in tabs.items():
        w.table(name, rows, JOB_BANDS, freeze_cols=3)
    w.table(
        "🏭 Companies",
        report.companies,
        [
            Band(
                "Company",
                "#305496",
                tuple(report.companies[0]) if report.companies else ("Company",),
            )
        ],
        freeze_cols=1,
    )
    w.table("🗑️ Rejected", rejected, [Band("Rejected", "#595959", REJECTED_COLUMNS)])
    w.key_values("🧾 Run Log", "Run log", run_info)
    w.close()
    return path
